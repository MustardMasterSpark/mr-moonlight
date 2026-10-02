"""Rebuild the two browsing sheets of the asset spreadsheet from its 'Asset Index' sheet.

    python Tools/pipeline/build_asset_browse_sheet.py            # master only
    python Tools/pipeline/build_asset_browse_sheet.py --mirror   # master + copy over Docs/asset-index/ASSETS-Index.xlsx

'Asset Index' stays the source of truth (sections, import process, Docs/asset-import-update-process.md).
It cannot be sorted or filtered directly, because section-divider rows sit in the middle of the table.
This tool therefore generates, and fully replaces on every run:

  Browse by Category - one flat table (no divider rows), filter buttons on every column, sorted by
                       category, then status (Owned first), then section, then id. A 'section' column
                       keeps what the divider rows said.
  By Category        - one row per category with live COUNTIFS formulas against the flat table
                       (owned / wishlist / total / in Mr. Moonlight).

Never edit those two sheets by hand: re-run this script after any change to 'Asset Index'.
Close the workbook in Excel first (openpyxl cannot save over Excel's lock).
"""
import copy
import os
import shutil
import sys

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

MASTER = r"C:\Users\calva\Documents\Asset Collection\03_documentation\ASSETS - Index 2026-09-16.xlsx"
MIRROR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "Docs", "asset-index", "ASSETS-Index.xlsx")
SRC, BROWSE, SUMMARY = "Asset Index", "Browse by Category", "By Category"

# Output column order -> source column header in 'Asset Index'. 'section' is derived, not copied.
COLUMNS = [
    "id", "asset name", "category", "status", "section", "priority", "price", "in Mr Moonlight?",
    "up to date?", "purpose in Mr Moonlight", "description", "notes", "collection url",
    "asset store / source", "documentation", "downloaded file", "publisher",
    "latest version (store)", "store updated",
]
WIDTH = {"section": 24}


def section_label(text):
    t = (text or "").upper()
    if t.startswith("OWNED — DOWNLOADED"):
        return "Owned"
    if t.startswith("OWNED / IN PROJECT"):
        return "Owned (added later)"
    if t.startswith("OWNED (FROM WISHLIST)"):
        return "Owned (not in project)"
    if t.startswith("WISHLIST"):
        for p in ("P1", "P2", "P3", "P4", "P5"):
            if " " + p + " " in t or t.split("—")[1].strip().startswith(p):
                return "Wishlist " + p
        return "Wishlist"
    return text or ""


def main():
    mirror = "--mirror" in sys.argv
    wb = openpyxl.load_workbook(MASTER)
    src = wb[SRC]
    hdr = {c.value: c.column for c in src[2] if c.value}
    missing = [h for h in COLUMNS if h != "section" and h not in hdr]
    if missing:
        sys.exit("Asset Index is missing columns: %s" % missing)

    # 1. read every asset row with its section, in sheet order
    rows, sec = [], ""
    for r in range(3, src.max_row + 1):
        ident = src.cell(r, hdr["id"]).value
        if ident and str(ident).startswith("AST-"):
            rows.append((r, section_label(sec)))
        else:
            label = src.cell(r, 2).value or src.cell(r, 1).value
            if label:
                sec = label
    sec_order = {}
    for _, s in rows:
        sec_order.setdefault(s, len(sec_order))

    def key(item):
        r, s = item
        status = src.cell(r, hdr["status"]).value or ""
        return ((src.cell(r, hdr["category"]).value or "~").lower(), 0 if status == "Owned" else 1,
                sec_order[s], str(src.cell(r, hdr["id"]).value))
    rows.sort(key=key)

    # 2. replace the generated sheets
    for name in (BROWSE, SUMMARY):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet(BROWSE, 0)
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines

    title = ws.cell(1, 2, "Browse by Category - generated from 'Asset Index'; use the filter buttons on row 2")
    title._style = copy.copy(src.cell(1, 2)._style)
    ws.row_dimensions[1].height = src.row_dimensions[1].height

    for j, name in enumerate(COLUMNS):
        col = 2 + j
        c = ws.cell(2, col, name)
        c._style = copy.copy(src.cell(2, hdr.get(name, hdr["status"]))._style)
        letter = get_column_letter(col)
        srcw = src.column_dimensions[get_column_letter(hdr[name])].width if name in hdr else None
        ws.column_dimensions[letter].width = WIDTH.get(name, srcw or 14)
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[2].height = src.row_dimensions[2].height

    for i, (r, s) in enumerate(rows):
        out = 3 + i
        for j, name in enumerate(COLUMNS):
            col = 2 + j
            if name == "section":
                c = ws.cell(out, col, s)
                c._style = copy.copy(src.cell(r, hdr["status"])._style)
                continue
            sc = src.cell(r, hdr[name])
            c = ws.cell(out, col, sc.value)
            c._style = copy.copy(sc._style)
            if sc.hyperlink is not None:
                link = copy.copy(sc.hyperlink)
                link.ref = c.coordinate
                c.hyperlink = link
        if src.row_dimensions[r].height:
            ws.row_dimensions[out].height = src.row_dimensions[r].height

    last = 2 + len(rows)
    ws.auto_filter.ref = "B2:%s%d" % (get_column_letter(1 + len(COLUMNS)), last)
    ws.freeze_panes = "D3"          # id + name stay visible, header stays visible

    # 3. live per-category summary
    sm = wb.create_sheet(SUMMARY, 1)
    sm.sheet_view.showGridLines = src.sheet_view.showGridLines
    heads = ["category", "owned", "wishlist", "total", "in Mr Moonlight (Yes)"]
    for j, h in enumerate(heads):
        c = sm.cell(2, 2 + j, h)
        c._style = copy.copy(src.cell(2, hdr["status"])._style)
    cats = sorted({src.cell(r, hdr["category"]).value or "(none)" for r, _ in rows}, key=str.lower)
    col = {n: get_column_letter(2 + COLUMNS.index(n)) for n in ("category", "status", "in Mr Moonlight?")}
    rng = lambda n: "'%s'!$%s$3:$%s$%d" % (BROWSE, col[n], col[n], last)
    for i, cat in enumerate(cats):
        r = 3 + i
        sm.cell(r, 2, cat)
        sm.cell(r, 3, '=COUNTIFS(%s,$B%d,%s,"Owned")' % (rng("category"), r, rng("status")))
        sm.cell(r, 4, '=COUNTIFS(%s,$B%d,%s,"Wishlist")' % (rng("category"), r, rng("status")))
        sm.cell(r, 5, "=C%d+D%d" % (r, r))
        sm.cell(r, 6, '=COUNTIFS(%s,$B%d,%s,"Yes*")' % (rng("category"), r, rng("in Mr Moonlight?")))
    tot = 3 + len(cats)
    sm.cell(tot, 2, "TOTAL").font = copy.copy(src.cell(2, 2).font)
    for k in range(3, 7):
        L = get_column_letter(k)
        sm.cell(tot, k, "=SUM(%s3:%s%d)" % (L, L, tot - 1))
    sm.column_dimensions["A"].width = 2
    sm.column_dimensions["B"].width = 36
    for L in "CDEF":
        sm.column_dimensions[L].width = 14
    sm.column_dimensions["F"].width = 22
    sm.auto_filter.ref = "B2:F%d" % (tot - 1)
    sm.freeze_panes = "C3"

    wb.active = 0
    for w in wb.worksheets:
        w.sheet_view.tabSelected = (w.title == BROWSE)
    wb.save(MASTER)
    print("wrote %d assets, %d categories -> %s" % (len(rows), len(cats), MASTER))
    if mirror:
        shutil.copyfile(MASTER, os.path.normpath(MIRROR))
        print("mirror updated:", os.path.normpath(MIRROR))


if __name__ == "__main__":
    main()
